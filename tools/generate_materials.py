from pathlib import Path
import re
import textwrap

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, PageBreak, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
OUTP = ROOT / "PPTs"
OUTPDF = ROOT / "PDFs"
AS = ROOT / "docs" / "assets"
for directory in (OUTP, OUTPDF, AS):
    directory.mkdir(parents=True, exist_ok=True)

# Rebuild generated material from scratch so stale/broken files cannot remain.
for directory, patterns in ((OUTP, ("*.pptx", "*_Preview.pdf")), (OUTPDF, ("*.pdf",)), (AS, ("*.png",))):
    for pattern in patterns:
        for old in directory.glob(pattern):
            old.unlink()

for old in (ROOT / "BUILD_TRIGGER.txt", ROOT / "tools" / "fsd2_source_bundle.b64"):
    if old.exists():
        old.unlink()
for old in (ROOT / "tools").glob("fsd2_bundle_*.b64"):
    old.unlink()

NAVY = RGBColor(27, 43, 75)
BLUE = RGBColor(48, 93, 156)
LIGHT_BLUE = RGBColor(232, 241, 252)
DARK = RGBColor(35, 39, 47)
GREY = RGBColor(102, 110, 122)
LIGHT_GREY = RGBColor(244, 246, 249)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(42, 122, 89)
ORANGE = RGBColor(214, 126, 45)
BORDER = RGBColor(220, 225, 232)


def read_text(path):
    return path.read_text(encoding="utf-8")


def title_tasks(readme):
    title = "FSD2 Lab Experiment"
    tasks = []
    for line in read_text(readme).splitlines():
        if line.startswith("# "):
            title = line[2:].strip()
        elif line.startswith("- "):
            tasks.append(line[2:].strip())
    return title, tasks


def load_font(size=34):
    for name in ("DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            pass
    return ImageFont.load_default()


def make_diagram(path, title, mode):
    width, height = 1600, 900
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = load_font(40 if len(title) > 65 else 46)
    draw.text((70, 45), title, font=title_font, fill=(27, 43, 75))
    if mode == "workflow":
        boxes = [("1. Input", "Browser / terminal"), ("2. Process", "Express / React"), ("3. Data", "State / database"), ("4. Output", "Screen / API")]
        footer = "Workflow: input → processing → data/state → output"
    else:
        boxes = [("1. Start", "Run the program"), ("2. Enter", "Give sample input"), ("3. Execute", "Logic runs"), ("4. Observe", "See the result")]
        footer = "Runtime: start → enter → execute → observe"
    xs = [55, 430, 805, 1180]
    for i, (heading, detail) in enumerate(boxes):
        x = xs[i]
        draw.rounded_rectangle((x, 285, x + 315, 525), radius=24, fill=(232, 241, 252), outline=(48, 93, 156), width=5)
        draw.text((x + 25, 330), heading, font=load_font(31), fill=(27, 43, 75))
        draw.text((x + 25, 400), detail, font=load_font(24), fill=(55, 62, 72))
        if i < 3:
            draw.line((x + 315, 405, x + 365, 405), fill=(48, 93, 156), width=8)
            draw.polygon([(x + 365, 405), (x + 345, 392), (x + 345, 418)], fill=(48, 93, 156))
    draw.text((70, 720), footer, font=load_font(27), fill=(102, 110, 122))
    image.save(path)


def collect_code_files(folder):
    result = []
    for path in folder.rglob("*"):
        if path.is_file() and path.suffix.lower() in {".js", ".jsx", ".ejs", ".html", ".css", ".scss"}:
            if "node_modules" not in path.parts and "dist" not in path.parts:
                result.append(path)
    return sorted(result)


def get_code_chunks(folder, max_lines=10, max_chunks=2):
    chunks = []
    for code_file in collect_code_files(folder)[:5]:
        lines = [line.rstrip().replace("\t", "  ") for line in read_text(code_file).splitlines()]
        if not any(line.strip() for line in lines):
            continue
        for start in range(0, min(len(lines), 30), max_lines):
            part = lines[start:start + max_lines]
            if part:
                chunks.append((code_file.name, start + 1, part))
            if len(chunks) >= max_chunks:
                return chunks
    return chunks


def add_bg(slide, color=WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    slide.shapes._spTree.remove(shape._element)
    slide.shapes._spTree.insert(2, shape._element)


def add_text(slide, text, x, y, w, h, size=20, color=DARK, bold=False, font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.05)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def add_header(slide, title, subtitle=None):
    title_size = 27 if len(title) < 58 else 23
    add_text(slide, title, 0.65, 0.30, 11.95, 0.68, size=title_size, color=NAVY, bold=True, valign=MSO_ANCHOR.MIDDLE)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.04), Inches(1.25), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()
    if subtitle:
        add_text(slide, subtitle, 2.05, 0.39, 10.2, 0.40, size=13, color=GREY)


def wrap_for_slide(text, width=72):
    return textwrap.wrap(text, width=width, break_long_words=False, break_on_hyphens=False) or [""]


def add_bullets(slide, items, y=1.55, font_size=18):
    for item in items:
        wrapped = wrap_for_slide(item)
        height = 0.42 + 0.28 * (len(wrapped) - 1)
        add_text(slide, "•", 0.85, y, 0.32, 0.4, size=22, color=BLUE, bold=True)
        add_text(slide, "\n".join(wrapped), 1.22, y - 0.01, 11.2, height, size=font_size, color=DARK)
        y += max(0.70, height + 0.18)


def add_title_slide(prs, title, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, LIGHT_BLUE)
    add_text(slide, f"FSD2 LAB • EXPERIMENT {number}", 0.8, 1.0, 11.7, 0.45, size=18, color=BLUE, bold=True)
    title_size = 38 if len(title) < 55 else 32 if len(title) < 80 else 27
    add_text(slide, title, 0.8, 1.78, 11.7, 1.75, size=title_size, color=NAVY, bold=True, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, "Simple explanation • Clean code • Runtime pathway • Viva preparation", 0.8, 3.82, 11.4, 0.55, size=19, color=GREY)
    add_text(slide, "FSD2 Lab Materials", 0.8, 6.35, 4.0, 0.4, size=16, color=BLUE, bold=True)


def add_task_slides(prs, tasks):
    tasks = tasks or ["See the experiment README for the exact syllabus task."]
    for page, start in enumerate(range(0, len(tasks), 5), start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_bg(slide)
        add_header(slide, "Syllabus tasks", f"Part {page}" if len(tasks) > 5 else None)
        add_bullets(slide, tasks[start:start + 5], y=1.55, font_size=18)


def add_code_slide(prs, index, filename, start_line, lines):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_header(slide, f"Code pathway {index}", f"{filename} • lines {start_line}–{start_line + len(lines) - 1}")
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(1.42), Inches(12.0), Inches(5.55))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(248, 250, 252)
    box.line.color.rgb = BORDER
    visual = []
    for offset, line in enumerate(lines):
        pieces = textwrap.wrap(line, width=94, subsequent_indent="    ", break_long_words=False, break_on_hyphens=False) or [""]
        visual.append((start_line + offset, pieces))
    code_lines = []
    for number, pieces in visual:
        code_lines.append(f"{number:>3}  {pieces[0]}")
        code_lines.extend(f"     {piece}" for piece in pieces[1:])
    add_text(slide, "\n".join(code_lines), 0.90, 1.68, 11.5, 5.05, size=12.5, color=DARK, font="DejaVu Sans Mono")


def add_image_slide(prs, title, image_path, caption):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_header(slide, title)
    slide.shapes.add_picture(str(image_path), Inches(0.65), Inches(1.40), width=Inches(12.0), height=Inches(5.55))
    add_text(slide, caption, 0.85, 7.04, 11.6, 0.25, size=12, color=GREY, align=PP_ALIGN.CENTER)


def add_card(slide, x, y, w, h, title, body, accent):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = LIGHT_GREY
    card.line.color.rgb = BORDER
    add_text(slide, title, x + 0.22, y + 0.18, w - 0.44, 0.4, size=18, color=accent, bold=True)
    add_text(slide, body, x + 0.22, y + 0.68, w - 0.44, h - 0.88, size=15.5, color=DARK)


def add_result_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_header(slide, "Expected result", title)
    add_card(slide, 0.75, 1.55, 3.65, 2.1, "INPUT", "Enter the sample values or open the required page / route.", BLUE)
    add_card(slide, 4.85, 1.55, 3.65, 2.1, "PROCESS", "The Express route, React component, hook, or MongoDB command performs the task.", GREEN)
    add_card(slide, 8.95, 1.55, 3.65, 2.1, "OUTPUT", "Check the browser, terminal, API response, or database result.", ORANGE)
    add_text(slide, "Verification checklist", 0.8, 4.15, 4.0, 0.45, size=22, color=NAVY, bold=True)
    add_bullets(slide, ["No syntax errors appear in the terminal.", "The requested feature responds to the input.", "The displayed / returned value matches the operation."], y=4.78, font_size=17)


def add_viva_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_header(slide, "Viva preparation")
    add_bullets(slide, ["What is the purpose of this experiment?", "Which input is required?", "Which line or command performs the main operation?", "What output should you observe?", "What happens when the input changes?"], y=1.55, font_size=18)


def make_preview_pdf(path, title, number, tasks, workflow, runtime, chunks):
    page_w, page_h = landscape((13.333 * inch, 7.5 * inch))
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleClean", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=25, leading=29, textColor=colors.HexColor("#1B2B4B"), spaceAfter=12)
    h_style = ParagraphStyle("HClean", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=19, leading=23, textColor=colors.HexColor("#1B2B4B"), spaceAfter=10)
    body = ParagraphStyle("BodyClean", parent=styles["BodyText"], fontName="Helvetica", fontSize=13, leading=17, textColor=colors.HexColor("#23272F"))
    code = ParagraphStyle("CodeClean", parent=body, fontName="Courier", fontSize=9.0, leading=11.3)
    doc = SimpleDocTemplate(str(path), pagesize=(page_w, page_h), leftMargin=0.65 * inch, rightMargin=0.65 * inch, topMargin=0.48 * inch, bottomMargin=0.45 * inch, title=title)
    story = [Paragraph(f"FSD2 LAB • EXPERIMENT {number}", body), Spacer(1, 6), Paragraph(title, title_style), Paragraph("GitHub-readable preview of the presentation. Download the matching .pptx for the editable PowerPoint.", body), PageBreak()]
    story.append(Paragraph("1. Syllabus tasks", h_style))
    data = [[Paragraph(f"• {t}", body)] for t in tasks[:10]]
    table = Table(data, colWidths=[11.95 * inch])
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F4F6F9")), ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#DCE1E8")), ("INNERGRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E6EAF0")), ("LEFTPADDING", (0, 0), (-1, -1), 10), ("RIGHTPADDING", (0, 0), (-1, -1), 10), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    story += [table, PageBreak(), Paragraph("2. Workflow", h_style), RLImage(str(workflow), width=11.95 * inch, height=5.85 * inch), PageBreak()]
    for i, (filename, start, lines) in enumerate(chunks, start=1):
        story.append(Paragraph(f"3.{i}. Code pathway — {filename}", h_style))
        escaped = "<br/>".join(f"{start + j:>3}  {line.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')}" for j, line in enumerate(lines))
        story.append(Table([[Paragraph(escaped, code)]], colWidths=[11.95 * inch], style=TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")), ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#D2D8E0")), ("LEFTPADDING", (0, 0), (-1, -1), 12), ("RIGHTPADDING", (0, 0), (-1, -1), 12), ("TOPPADDING", (0, 0), (-1, -1), 10), ("BOTTOMPADDING", (0, 0), (-1, -1), 10)])))
        story.append(PageBreak())
    story += [Paragraph("4. Runtime walkthrough", h_style), RLImage(str(runtime), width=11.95 * inch, height=5.85 * inch), PageBreak(), Paragraph("5. Run and observe", h_style)]
    for item in ["Start the server / Vite app or open the MongoDB shell as described in the experiment README.", "Enter the sample input or call the required route.", "Watch the terminal for logs and errors.", "Verify the browser, API response, or database result."]:
        story += [Paragraph(f"• {item}", body), Spacer(1, 5)]
    story += [Spacer(1, 10), Paragraph("6. Viva questions", h_style)]
    for item in ["What is the purpose of this experiment?", "Which input is required?", "Which line or command performs the main operation?", "What output should you observe?", "What happens when the input changes?"]:
        story.append(Paragraph(f"• {item}", body))
    doc.build(story)


def build(folder):
    title, tasks = title_tasks(folder / "README.md")
    number = folder.name[:2]
    slug = re.sub(r"[^A-Za-z0-9]+", "_", title).strip("_")
    workflow = AS / f"{number}_workflow.png"
    runtime = AS / f"{number}_runtime.png"
    make_diagram(workflow, title, "workflow")
    make_diagram(runtime, "Runtime walkthrough — " + title, "runtime")
    chunks = get_code_chunks(folder)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    add_title_slide(prs, title, number)
    add_task_slides(prs, tasks)
    add_image_slide(prs, "How it works", workflow, "Input → processing → data/state → output")
    for i, chunk in enumerate(chunks, start=1):
        add_code_slide(prs, i, *chunk)
    add_image_slide(prs, "Runtime walkthrough", runtime, "Start → enter input → execute → observe output")
    add_result_slide(prs, title)
    add_viva_slide(prs)

    ppt_path = OUTP / f"{number}_{slug}.pptx"
    prs.save(ppt_path)
    preview_path = OUTP / f"{number}_{slug}_Preview.pdf"
    make_preview_pdf(preview_path, title, number, tasks, workflow, runtime, chunks)
    runtime_pdf = OUTPDF / f"{number}_{slug}_Runtime.pdf"
    make_preview_pdf(runtime_pdf, title, number, tasks, workflow, runtime, chunks)
    return number, title, slug


rows = []
for folder in sorted(ROOT.iterdir()):
    if folder.is_dir() and re.match(r"^\d\d-", folder.name) and (folder / "README.md").exists():
        rows.append(build(folder))

index_lines = [
    "# FSD2 Lab Materials",
    "",
    "Clean 16:9 presentations, GitHub-readable PDF previews, runtime guides and workflow visuals.",
    "",
    "> **How to view in GitHub:** open the `*_Preview.pdf` file inside `PPTs`. GitHub does not render `.pptx` slides in its file viewer. The `.pptx` is the editable PowerPoint version.",
    "",
]
for number, title, slug in rows:
    index_lines += [
        f"## {number}. {title}",
        f"- [Editable PowerPoint](PPTs/{number}_{slug}.pptx)",
        f"- [Open presentation in GitHub](PPTs/{number}_{slug}_Preview.pdf)",
        f"- [Runtime PDF](PDFs/{number}_{slug}_Runtime.pdf)",
        f"- [Workflow visual](docs/assets/{number}_workflow.png)",
        f"- [Runtime visual](docs/assets/{number}_runtime.png)",
        "",
    ]
(ROOT / "MATERIALS_INDEX.md").write_text("\n".join(index_lines), encoding="utf-8")

(OUTP / "README.md").write_text("""# PPTs — FSD2 Lab Presentations\n\nEach experiment has:\n\n- `*.pptx` — editable PowerPoint file.\n- `*_Preview.pdf` — GitHub-readable version of the same presentation.\n\nThe presentations use a consistent 16:9 layout, readable text, short points, clean code blocks, workflow visuals, runtime walkthroughs, expected-result checks and viva questions.\n\n**To view a presentation directly in GitHub, open its `*_Preview.pdf`.**\n\nSee [`MATERIALS_INDEX.md`](../MATERIALS_INDEX.md) for the complete list.\n""", encoding="utf-8")

print(f"Generated clean FSD2 presentations, GitHub previews and runtime guides for {len(rows)} experiments.")
